"""Leitura e gravação da planilha.

É a ÚNICA camada que fala com a base de dados. Existem dois "motores":

* GoogleSheets: o modo real, pela conta de serviço (gspread).
* Local: um arquivo .xlsx habilitado somente por ``GA_FORCAR_LOCAL=1``, para
  desenvolvimento e testes. A ausência de credenciais nunca causa fallback.

As páginas nunca chamam o gspread diretamente: usam as funções públicas
deste módulo (ler, ler_fresco, inserir, atualizar, excluir...).
"""
from __future__ import annotations

import functools
import os
import re
import threading
import uuid
from pathlib import Path

import pandas as pd
import streamlit as st

from servicos.esquema import ABAS_COM_ID, COLUNAS, LISTAS_PADRAO, ABA_LISTAS

MSG_FALHA_CONEXAO = (
    "Não foi possível acessar a planilha. Tente de novo em alguns segundos."
)


class ErroPlanilha(Exception):
    """Erro com mensagem em português, pronta para mostrar ao usuário."""


def _traduz_erros(func):
    """Transforma qualquer falha técnica em ErroPlanilha com mensagem clara."""

    @functools.wraps(func)
    def wrapper(*args, **kwargs):
        try:
            return func(*args, **kwargs)
        except ErroPlanilha:
            raise
        except Exception as exc:  # noqa: BLE001 - queremos capturar tudo aqui
            texto = str(exc)
            if "PERMISSION_DENIED" in texto or "403" in texto:
                raise ErroPlanilha(
                    "A conta de serviço não tem permissão na planilha. "
                    "Compartilhe a planilha como Editor com o e-mail da conta de serviço."
                ) from exc
            if "NOT_FOUND" in texto or "404" in texto or "SpreadsheetNotFound" in type(exc).__name__:
                raise ErroPlanilha(
                    "Planilha não encontrada. Confira o link em [planilha] url nos Secrets."
                ) from exc
            if "WorksheetNotFound" in type(exc).__name__:
                raise ErroPlanilha(
                    f"A aba {texto} não existe na planilha. Rode o script "
                    "scripts/criar_planilha_modelo.py ou crie a aba com esse nome."
                ) from exc
            raise ErroPlanilha(MSG_FALHA_CONEXAO) from exc

    return wrapper


# --------------------------------------------------------------------------
# Motores
# --------------------------------------------------------------------------
def _para_sheets(v):
    """Grava em modo RAW (sem a planilha "interpretar" o texto): datas ISO viram
    número de série do Sheets e o texto é gravado como texto. Assim um "150.00"
    ou "1/2" nunca é confundido com outra coisa por causa da localidade."""
    if isinstance(v, str):
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}( \d{2}:\d{2}(:\d{2})?)?", v):
            dias = (pd.Timestamp(v) - pd.Timestamp("1899-12-30")) / pd.Timedelta(days=1)
            return int(dias) if float(dias).is_integer() else round(float(dias), 6)
    return v


class MotorGoogleSheets:
    nome = "Google Sheets"
    _OPCAO = "RAW"

    def __init__(self, credenciais: dict, url_ou_id: str):
        import gspread

        cliente = gspread.service_account_from_dict(dict(credenciais))
        if url_ou_id.startswith("http"):
            self._planilha = cliente.open_by_url(url_ou_id)
        else:
            self._planilha = cliente.open_by_key(url_ou_id)
        self._abas = {}

    def _aba(self, nome):
        if nome not in self._abas:
            try:
                self._abas[nome] = self._planilha.worksheet(nome)
            except Exception as exc:
                if "WorksheetNotFound" not in type(exc).__name__ or nome not in COLUNAS:
                    raise
                # Migração aditiva: cria novas abas do aplicativo sem alterar dados existentes.
                ws = self._planilha.add_worksheet(title=nome, rows=1000, cols=max(10, len(COLUNAS[nome])))
                ws.append_row(COLUNAS[nome], value_input_option=self._OPCAO)
                self._abas[nome] = ws
        return self._abas[nome]

    def ler(self, aba):
        # UNFORMATTED_VALUE: números chegam como número e datas como número de
        # série, independentemente da localidade da planilha.
        return self._aba(aba).get_all_values(value_render_option="UNFORMATTED_VALUE")

    def anexar(self, aba, linhas):
        self._aba(aba).append_rows(
            [[_para_sheets(v) for v in linha] for linha in linhas],
            value_input_option=self._OPCAO, table_range="A1", insert_data_option="INSERT_ROWS",
        )

    def atualizar_linhas(self, aba, linhas: dict[int, list]):
        if not linhas:
            return
        dados = [{"range": f"A{n}", "values": [[_para_sheets(v) for v in valores]]}
                 for n, valores in linhas.items()]
        self._aba(aba).batch_update(dados, value_input_option=self._OPCAO)

    def atualizar_celulas(self, aba, celulas: dict[str, object]):
        if not celulas:
            return
        dados = [{"range": rng, "values": [[_para_sheets(v)]]} for rng, v in celulas.items()]
        self._aba(aba).batch_update(dados, value_input_option=self._OPCAO)

    def excluir_linha(self, aba, numero):
        self._aba(aba).delete_rows(numero)


class MotorLocal:
    """Planilha simulada num .xlsx, com as mesmas abas e colunas."""

    nome = "modo local (demonstração)"
    _trava = threading.Lock()

    def __init__(self, caminho: Path):
        self.caminho = Path(caminho)
        self.caminho.parent.mkdir(parents=True, exist_ok=True)
        if not self.caminho.exists():
            self._criar_vazia()

    def _criar_vazia(self):
        from openpyxl import Workbook

        wb = Workbook()
        wb.remove(wb.active)
        for aba, colunas in COLUNAS.items():
            ws = wb.create_sheet(aba)
            ws.append(colunas)
        ws = wb[ABA_LISTAS]
        maior = max(len(v) for v in LISTAS_PADRAO.values())
        for i in range(maior):
            ws.append([LISTAS_PADRAO[c][i] if i < len(LISTAS_PADRAO[c]) else None
                       for c in COLUNAS[ABA_LISTAS]])
        wb.save(self.caminho)

    def _abrir(self):
        from openpyxl import load_workbook

        return load_workbook(self.caminho)

    def ler(self, aba):
        with self._trava:
            wb = self._abrir()
            if aba not in wb.sheetnames:
                ws = wb.create_sheet(aba)
                ws.append(COLUNAS[aba])
                wb.save(self.caminho)
            ws = wb[aba]
            return [["" if v is None else v for v in linha]
                    for linha in ws.iter_rows(values_only=True)]

    def anexar(self, aba, linhas):
        with self._trava:
            wb = self._abrir()
            ws = wb[aba]
            # ignora linhas totalmente vazias no fim
            ultima = ws.max_row
            while ultima > 1 and all(c.value in (None, "") for c in ws[ultima]):
                ultima -= 1
            for i, linha in enumerate(linhas, start=1):
                for j, v in enumerate(linha, start=1):
                    ws.cell(row=ultima + i, column=j, value=_valor_local(v))
            wb.save(self.caminho)

    def atualizar_linhas(self, aba, linhas):
        with self._trava:
            wb = self._abrir()
            ws = wb[aba]
            for n, valores in linhas.items():
                for j, v in enumerate(valores, start=1):
                    ws.cell(row=n, column=j, value=_valor_local(v))
            wb.save(self.caminho)

    def atualizar_celulas(self, aba, celulas):
        with self._trava:
            wb = self._abrir()
            ws = wb[aba]
            for rng, v in celulas.items():
                ws[rng] = _valor_local(v)
            wb.save(self.caminho)

    def excluir_linha(self, aba, numero):
        with self._trava:
            wb = self._abrir()
            wb[aba].delete_rows(numero)
            wb.save(self.caminho)


def _valor_local(v):
    """No modo local, grava datas ISO como data de verdade (igual ao Sheets)."""
    if isinstance(v, str):
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
            return pd.Timestamp(v).to_pydatetime().date()
        if re.fullmatch(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}", v):
            return pd.Timestamp(v).to_pydatetime()
    return v


# --------------------------------------------------------------------------
# Escolha do motor
# --------------------------------------------------------------------------
def _segredo(secao, chave=None, padrao=None):
    try:
        if secao not in st.secrets:
            return padrao
        valor = st.secrets[secao]
        return valor if chave is None else valor.get(chave, padrao)
    except Exception:  # sem secrets.toml
        return padrao


CAMINHO_LOCAL = Path(
    os.environ.get("GA_PLANILHA_LOCAL",
                   Path(__file__).resolve().parent.parent / "dados_locais" / "planilha_local.xlsx")
)


@st.cache_resource(show_spinner=False)
def _motor_cache(caminho_local: str):
    credenciais = _segredo("gcp_service_account")
    url = _segredo("planilha", "url")
    if os.environ.get("GA_FORCAR_LOCAL") == "1":
        return MotorLocal(Path(caminho_local))
    if not credenciais or not url:
        raise ErroPlanilha(
            "O Google Sheets não está configurado. Defina [planilha] url e "
            "[gcp_service_account] nos Secrets."
        )
    return _traduz_erros(MotorGoogleSheets)(credenciais, url)


def motor():
    return _motor_cache(str(CAMINHO_LOCAL))


def modo_local() -> bool:
    return os.environ.get("GA_FORCAR_LOCAL") == "1"


# --------------------------------------------------------------------------
# Leitura
# --------------------------------------------------------------------------
PREFIXO_ID = {
    "Assinaturas": "A", "Funcionários": "F", "Projetos": "P",
    "Uso por projeto": "U", "Equipe do projeto": "E", "Usuários": "L",
}


def novo_id(aba: str = "") -> str:
    """ID curto com letra na frente (o Sheets não o confunde com número)."""
    return f"{PREFIXO_ID.get(aba, 'X')}-{uuid.uuid4().hex[:6].upper()}"


def _para_df(aba, valores) -> pd.DataFrame:
    colunas = COLUNAS[aba]
    linhas = []
    for numero, linha in enumerate(valores[1:], start=2):
        linha = list(linha) + [""] * (len(colunas) - len(linha))
        linha = linha[: len(colunas)]
        if all(str(v).strip() == "" for v in linha):
            continue
        linhas.append(dict(zip(colunas, linha)) | {"_linha": numero})
    df = pd.DataFrame(linhas, columns=colunas + ["_linha"])
    for col in colunas:
        df[col] = df[col].astype(object)
    return df


@_traduz_erros
def ler_fresco(aba: str) -> pd.DataFrame:
    """Lê a aba direto da fonte, sem cache. Preenche IDs que faltarem."""
    m = motor()
    df = _para_df(aba, m.ler(aba))
    if aba in ABAS_COM_ID and not df.empty:
        sem_id = df["ID"].astype(str).str.strip() == ""
        if sem_id.any():
            # Linhas incluídas direto na planilha: o app gera o ID.
            celulas = {}
            for idx in df.index[sem_id]:
                novo = novo_id(aba)
                df.at[idx, "ID"] = novo
                celulas[f"A{df.at[idx, '_linha']}"] = novo
            m.atualizar_celulas(aba, celulas)
    if "ID" in df.columns:
        df["ID"] = df["ID"].astype(str)
    return df


@st.cache_data(ttl=60, show_spinner="Lendo a planilha...")
def _ler_cache(aba: str, _chave_motor: str) -> pd.DataFrame:
    return ler_fresco(aba)


def ler(aba: str) -> pd.DataFrame:
    """Lê a aba com cache de 60 segundos."""
    return _ler_cache(aba, str(id(motor()))).copy()


def limpar_cache():
    _ler_cache.clear()


def listas() -> dict[str, list[str]]:
    """Opções das listas suspensas (aba Listas), com os padrões como reserva."""
    df = ler(ABA_LISTAS)
    resultado = {}
    for col in COLUNAS[ABA_LISTAS]:
        valores = [str(v).strip() for v in df[col].tolist() if str(v).strip()] if not df.empty else []
        resultado[col] = valores or list(LISTAS_PADRAO[col])
    return resultado


# --------------------------------------------------------------------------
# Gravação
# --------------------------------------------------------------------------
def _linha_de(aba, registro: dict) -> list:
    return [_celula(registro.get(col, "")) for col in COLUNAS[aba]]


def _celula(v):
    if v is None:
        return ""
    if isinstance(v, float) and pd.isna(v):
        return ""
    if isinstance(v, pd.Timestamp):
        return v.strftime("%Y-%m-%d %H:%M:%S") if (v.hour or v.minute or v.second) else v.strftime("%Y-%m-%d")
    if hasattr(v, "isoformat") and not isinstance(v, str):
        return v.isoformat(sep=" ") if hasattr(v, "hour") else v.isoformat()
    return v


@_traduz_erros
def numero_da_linha(aba: str, id_: str) -> int | None:
    """Acha a linha pelo ID, relendo a aba (alguém pode ter ordenado/inserido)."""
    df = ler_fresco(aba)
    achou = df[df["ID"].astype(str) == str(id_)]
    return None if achou.empty else int(achou["_linha"].iloc[0])


@_traduz_erros
def ler_registro(aba: str, id_: str) -> dict | None:
    df = ler_fresco(aba)
    achou = df[df["ID"].astype(str) == str(id_)]
    return None if achou.empty else achou.iloc[0].to_dict()


@_traduz_erros
def inserir(aba: str, registros: list[dict]):
    motor().anexar(aba, [_linha_de(aba, r) for r in registros])


@_traduz_erros
def atualizar_varios(aba: str, registros: dict[str, dict]):
    """registros = {ID: registro completo}. Uma única chamada à planilha."""
    if not registros:
        return
    df = ler_fresco(aba)
    por_id = dict(zip(df["ID"].astype(str), df["_linha"]))
    linhas = {}
    for id_, reg in registros.items():
        if str(id_) not in por_id:
            raise ErroPlanilha("Um dos registros não existe mais na planilha. Clique em Atualizar.")
        linhas[int(por_id[str(id_)])] = _linha_de(aba, reg)
    motor().atualizar_linhas(aba, linhas)


def atualizar(aba: str, id_: str, registro: dict):
    atualizar_varios(aba, {id_: registro})


def _letra_coluna(numero: int) -> str:
    letras = ""
    while numero:
        numero, resto = divmod(numero - 1, 26)
        letras = chr(65 + resto) + letras
    return letras


@_traduz_erros
def atualizar_campos(aba: str, id_: str, campos: dict[str, object]):
    """Atualiza somente as células informadas, relocalizando a linha pelo ID."""
    desconhecidos = set(campos) - set(COLUNAS[aba])
    if desconhecidos:
        raise ErroPlanilha(f"Campo(s) desconhecido(s): {', '.join(sorted(desconhecidos))}.")
    numero = numero_da_linha(aba, id_)
    if numero is None:
        raise ErroPlanilha("Esse registro não existe mais na planilha. Clique em Atualizar.")
    celulas = {
        f"{_letra_coluna(COLUNAS[aba].index(campo) + 1)}{numero}": _celula(valor)
        for campo, valor in campos.items()
    }
    motor().atualizar_celulas(aba, celulas)


@_traduz_erros
def excluir(aba: str, id_: str):
    n = numero_da_linha(aba, id_)
    if n is None:
        raise ErroPlanilha("Esse registro já foi excluído da planilha.")
    motor().excluir_linha(aba, n)
