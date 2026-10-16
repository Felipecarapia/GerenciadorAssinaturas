"""Geração dos 5 relatórios (RF-27 a RF-30) em PDF (fpdf2) e Excel (openpyxl)."""
from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date, datetime

import pandas as pd
from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from servicos import calculos as calc
from servicos import graficos

AZUL = (37, 99, 235)
ROXO = (99, 102, 241)
ESCURO = (15, 23, 42)
CINZA = (100, 116, 139)
CLARO = (241, 245, 249)
COR_NIVEL = {"vermelho": (220, 38, 38), "amarelo": (217, 119, 6), "verde": (5, 150, 105), "nenhum": CINZA}
NOME_NIVEL = {"vermelho": "Urgente", "amarelo": "Atenção", "verde": "Em dia", "nenhum": "-"}


@dataclass
class Dados:
    assinaturas: pd.DataFrame  # já passado por calculos.preparar_assinaturas
    funcionarios: pd.DataFrame
    historico: pd.DataFrame
    agora: datetime
    empresa: str = "RECÔNCAVO"

    @property
    def hoje(self) -> date:
        return self.agora.date()


# --------------------------------------------------------------------------
# Imagens dos gráficos para o PDF
# --------------------------------------------------------------------------
def _png_plotly(fig, largura=900, altura=420) -> bytes | None:
    """Kaleido (precisa do Chrome/Chromium). Retorna None se não estiver disponível."""
    try:
        fig = fig.to_dict()
        import plotly.graph_objects as go
        f = go.Figure(fig)
        f.update_layout(paper_bgcolor="white", plot_bgcolor="white",
                        font=dict(color="#0f172a"), height=altura, width=largura)
        return f.to_image(format="png", scale=2)
    except Exception:
        return None


def _png_matplotlib(tipo: str, df: pd.DataFrame) -> bytes:
    """Reserva para quando o Kaleido não roda (ex.: servidor sem Chromium)."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(9, 4.2), dpi=160)
    plt.rcParams["font.family"] = "DejaVu Sans"
    if df.empty:
        ax.text(0.5, 0.5, "Sem dados", ha="center", va="center", color="#94a3b8", fontsize=14)
        ax.axis("off")
    elif tipo == "donut":
        total = df["Gasto mensal"].sum()
        fatias, _ = ax.pie(df["Gasto mensal"], colors=(graficos.PALETA * 3)[: len(df)], startangle=90,
                           counterclock=False, wedgeprops=dict(width=0.42, edgecolor="white"))
        ax.set_aspect("equal")
        ax.text(0, 0, f"{calc.formatar_reais(total)}\npor mês", ha="center", va="center",
                fontsize=11, fontweight="bold", color="#0f172a")
        ax.legend(fatias, [f"{c}  {calc.formatar_reais(v)} ({v / total:.0%})"
                           for c, v in zip(df["Categoria"], df["Gasto mensal"])],
                  loc="center left", bbox_to_anchor=(1.0, 0.5), frameon=False, fontsize=9)
    elif tipo == "linha":
        ax.plot(df["Mês"], df["Gasto mensal"], color="#6366f1", linewidth=2.5, marker="o")
        ax.fill_between(range(len(df)), df["Gasto mensal"], color="#6366f1", alpha=0.12)
        ax.set_ylim(bottom=0)
        ax.tick_params(axis="x", labelsize=8, rotation=45)
        ax.yaxis.set_major_formatter(lambda v, _: calc.formatar_reais(v))
        for s in ("top", "right"):
            ax.spines[s].set_visible(False)
    elif tipo == "top5":
        d = df.iloc[::-1]
        barras = ax.barh(d["Nome"], d["Gasto mensal"], color="#3b82f6")
        ax.bar_label(barras, labels=[calc.formatar_reais(v) for v in d["Gasto mensal"]], padding=4, fontsize=9)
        ax.set_xlim(0, d["Gasto mensal"].max() * 1.35)
        ax.xaxis.set_visible(False)
        for s in ("top", "right", "bottom"):
            ax.spines[s].set_visible(False)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


def imagem_grafico(tipo: str, df: pd.DataFrame) -> bytes:
    construtor = {"donut": graficos.donut_categoria, "linha": graficos.linha_evolucao,
                  "top5": graficos.barras_top5}[tipo]
    png = _png_plotly(construtor(df)) if not df.empty else None
    return png or _png_matplotlib(tipo, df)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def _t(texto) -> str:
    """As fontes padrão do PDF usam Latin-1: troca o que não cabe."""
    s = "" if texto is None else str(texto)
    s = (s.replace("—", "-").replace("–", "-").replace("“", '"').replace("”", '"')
          .replace("’", "'").replace("•", "-").replace("…", "...").replace("·", "-"))
    return s.encode("latin-1", "replace").decode("latin-1")


class PDF(FPDF):
    def __init__(self, titulo: str, subtitulo: str, gerado_em: datetime, retrato=True):
        super().__init__(orientation="P" if retrato else "L", unit="mm", format="A4")
        self.titulo, self.subtitulo, self.gerado_em = titulo, subtitulo, gerado_em
        self.set_auto_page_break(True, margin=16)
        self.set_margins(14, 14, 14)
        self.set_title(_t(titulo))
        self.set_author("RECÔNCAVO")
        self.add_page()

    def header(self):
        self.set_fill_color(*AZUL)
        self.rect(0, 0, self.w, 3, "F")
        self.set_fill_color(*ROXO)
        self.rect(self.w * 0.6, 0, self.w * 0.4, 3, "F")
        if self.page_no() == 1:
            self.set_y(12)
            self.set_font("Helvetica", "B", 18)
            self.set_text_color(*ESCURO)
            self.cell(0, 9, _t(self.titulo), new_x="LMARGIN", new_y="NEXT")
            self.set_font("Helvetica", "", 10)
            self.set_text_color(*CINZA)
            self.cell(0, 5, _t(f"{self.subtitulo}  |  Gerado em {self.gerado_em:%d/%m/%Y %H:%M}"),
                      new_x="LMARGIN", new_y="NEXT")
            self.ln(5)
        else:
            self.set_y(9)
            self.set_font("Helvetica", "", 8)
            self.set_text_color(*CINZA)
            self.cell(0, 5, _t(self.titulo), new_x="LMARGIN", new_y="NEXT")
            self.ln(3)

    def footer(self):
        self.set_y(-11)
        self.set_font("Helvetica", "", 8)
        self.set_text_color(*CINZA)
        self.cell(0, 5, _t("RECÔNCAVO · Gerenciador de Assinaturas"), align="L")
        self.cell(0, 5, f"{self.page_no()}/{{nb}}", align="R")

    def secao(self, texto, espaco=40):
        """Título de seção; vai para a próxima página se não couber o conteúdo logo abaixo."""
        if self.get_y() + 12 + espaco > self.h - 16:
            self.add_page()
        self.ln(2)
        self.set_font("Helvetica", "B", 12)
        self.set_text_color(*ESCURO)
        self.cell(0, 7, _t(texto), new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(*AZUL)
        self.set_line_width(0.5)
        self.line(self.l_margin, self.get_y(), self.l_margin + 18, self.get_y())
        self.ln(3)

    def paragrafo(self, texto, tamanho=9):
        self.set_font("Helvetica", "", tamanho)
        self.set_text_color(*CINZA)
        self.multi_cell(0, 5, _t(texto), new_x="LMARGIN", new_y="NEXT")
        self.ln(1)

    def cards(self, itens: list[tuple[str, str]]):
        largura = (self.epw - 3 * (len(itens) - 1)) / len(itens)
        y = self.get_y()
        for i, (rotulo, valor) in enumerate(itens):
            x = self.l_margin + i * (largura + 3)
            self.set_fill_color(*CLARO)
            self.rect(x, y, largura, 20, "F", round_corners=True, corner_radius=2)
            self.set_xy(x + 3, y + 3)
            self.set_font("Helvetica", "", 7.5)
            self.set_text_color(*CINZA)
            self.cell(largura - 6, 4, _t(rotulo.upper()))
            self.set_xy(x + 3, y + 9)
            self.set_font("Helvetica", "B", 12.5)
            self.set_text_color(*ESCURO)
            self.cell(largura - 6, 7, _t(valor))
        self.set_xy(self.l_margin, y + 25)

    def imagem(self, png: bytes, largura=None):
        largura = largura or self.epw
        altura = largura * 420 / 900
        if self.get_y() + altura > self.h - 18:
            self.add_page()
        self.image(io.BytesIO(png), x=self.l_margin + (self.epw - largura) / 2, w=largura)
        self.ln(2)

    def tabela(self, df: pd.DataFrame, larguras=None, alinhar=None, cores_linha=None):
        if df.empty:
            self.paragrafo("Nenhum registro.")
            return
        self.set_font("Helvetica", "", 8.5)
        self.set_text_color(*ESCURO)
        self.set_draw_color(226, 232, 240)
        self.set_line_width(0.2)
        cab = FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=AZUL)
        with self.table(
            col_widths=larguras, text_align=alinhar or "LEFT", headings_style=cab,
            line_height=5.6, padding=1.4, cell_fill_color=(248, 250, 252), cell_fill_mode="ROWS",
            borders_layout="HORIZONTAL_LINES", first_row_as_headings=True,
        ) as tabela:
            linha = tabela.row()
            for col in df.columns:
                linha.cell(_t(col))
            for i, (_, r) in enumerate(df.iterrows()):
                linha = tabela.row()
                for j, v in enumerate(r):
                    estilo = None
                    if cores_linha and j == 0 and cores_linha[i]:
                        estilo = FontFace(emphasis="BOLD", color=cores_linha[i])
                    linha.cell(_t(v), style=estilo)
        self.ln(3)


def _bytes(pdf: PDF) -> bytes:
    return bytes(pdf.output())


def _tabela_assinaturas(df: pd.DataFrame) -> pd.DataFrame:
    return pd.DataFrame({
        "Assinatura": df["Nome"], "Categoria": df["Categoria"], "Plano": df["Plano"],
        "Valor": df["valor_num"].map(calc.formatar_reais),
        "Mensal": df["mensal"].map(calc.formatar_reais),
        "Vencimento": df["venc_data"].map(calc.formatar_data),
        "Status": df["Status"],
    })


def resumo_geral_pdf(d: Dados) -> bytes:
    """RF-27: assinaturas, gastos e colaboradores."""
    ind = calc.indicadores(d.assinaturas, d.funcionarios)
    pdf = PDF("Resumo Geral", "Assinaturas, gastos e colaboradores", d.agora)
    pdf.alias_nb_pages()
    pdf.cards([
        ("Assinaturas ativas", str(ind["total_ativas"])),
        ("Gasto mensal", calc.formatar_reais(ind["gasto_mensal"])),
        ("Gasto anual", calc.formatar_reais(ind["gasto_anual"])),
        ("Vencem em 30 dias", str(ind["vencendo_30"])),
        ("Colaboradores", str(ind["colaboradores"])),
    ])
    pdf.paragrafo("Regra do gasto: somam só as assinaturas ativas; planos anuais entram divididos "
                  "por 12 (trimestral por 3, semestral por 6); vitalícios não entram.", 8)

    pdf.secao("Gasto mensal por categoria", espaco=150 * 420 / 900)
    cat = calc.gasto_por_categoria(d.assinaturas)
    pdf.imagem(imagem_grafico("donut", cat), largura=150)

    pdf.secao("Evolução do gasto nos últimos 12 meses", espaco=170 * 420 / 900)
    pdf.imagem(imagem_grafico("linha", calc.evolucao_12_meses(
        d.assinaturas, d.historico, d.agora)), largura=170)

    pdf.secao("Top 5 mais caras (custo mensal)", espaco=160 * 420 / 900)
    pdf.imagem(imagem_grafico("top5", calc.top_mais_caras(d.assinaturas)), largura=160)

    pdf.secao("Assinaturas ativas")
    at = calc.ativas(d.assinaturas).sort_values("mensal", ascending=False)
    pdf.tabela(_tabela_assinaturas(at), larguras=(34, 26, 19, 23, 23, 22, 18),
               alinhar=("LEFT", "LEFT", "LEFT", "RIGHT", "RIGHT", "CENTER", "LEFT"))

    pdf.secao("Gastos por departamento")
    dep = calc.gasto_por_departamento(d.assinaturas, d.funcionarios)
    pdf.tabela(pd.DataFrame({
        "Departamento": dep["Departamento"], "Colaboradores": dep["Colaboradores"],
        "Gasto mensal": dep["Gasto mensal"].map(calc.formatar_reais),
        "Gasto anual": dep["Gasto anual"].map(calc.formatar_reais),
    }), alinhar=("LEFT", "CENTER", "RIGHT", "RIGHT"))
    return _bytes(pdf)


def a_vencer_pdf(d: Dados, dias: int) -> bytes:
    """RF-29: assinaturas a vencer em 30, 60 ou 90 dias (inclui as já vencidas)."""
    df = calc.a_vencer(d.assinaturas, dias)
    pdf = PDF(f"Assinaturas a vencer em {dias} dias", f"Até {calc.formatar_data(d.hoje + pd.Timedelta(days=dias))}",
              d.agora)
    pdf.alias_nb_pages()
    total = float(df["valor_num"].fillna(0).sum()) if not df.empty else 0.0
    vencidas = int((df["dias"].astype(int) < 0).sum()) if not df.empty else 0
    pdf.cards([("Assinaturas", str(len(df))), ("Já vencidas", str(vencidas)),
               ("Valor a renovar", calc.formatar_reais(total))])
    pdf.paragrafo("Urgente: vencida ou até 7 dias  |  Atenção: 8 a 30 dias  |  Em dia: acima de 30 dias", 8)
    if df.empty:
        pdf.paragrafo("Nenhuma assinatura vence nesse período.")
        return _bytes(pdf)
    usuarios = calc.usuarios_por_assinatura(d.funcionarios)
    tabela = pd.DataFrame({
        "Alerta": df["nivel"].map(NOME_NIVEL),
        "Assinatura": df["Nome"], "Plano": df["Plano"],
        "Valor": df["valor_num"].map(calc.formatar_reais),
        "Vencimento": df["venc_data"].map(calc.formatar_data),
        "Situação": df["texto_alerta"],
        "Usuários": df["Nome"].map(lambda n: len(usuarios.get(str(n), []))),
    })
    cores = [COR_NIVEL[n] for n in df["nivel"]]
    pdf.tabela(tabela, larguras=(18, 38, 20, 24, 22, 32, 18),
               alinhar=("LEFT", "LEFT", "LEFT", "RIGHT", "CENTER", "LEFT", "CENTER"),
               cores_linha=cores)
    return _bytes(pdf)


def colaboradores_pdf(d: Dados) -> bytes:
    """RF-30 (PDF): colaboradores por assinatura e gastos por departamento."""
    pdf = PDF("Colaboradores e Gastos por Departamento",
              "Quem usa cada assinatura e quanto cada área representa", d.agora, retrato=False)
    pdf.alias_nb_pages()
    pdf.paragrafo("O custo de cada assinatura é dividido igualmente entre os seus usuários.", 8)
    pdf.secao("Gastos por departamento")
    dep = calc.gasto_por_departamento(d.assinaturas, d.funcionarios)
    pdf.tabela(pd.DataFrame({
        "Departamento": dep["Departamento"], "Colaboradores": dep["Colaboradores"],
        "Gasto mensal": dep["Gasto mensal"].map(calc.formatar_reais),
        "Gasto anual": dep["Gasto anual"].map(calc.formatar_reais),
    }), alinhar=("LEFT", "CENTER", "RIGHT", "RIGHT"))
    pdf.secao("Colaboradores por assinatura")
    cpa = calc.colaboradores_por_assinatura(d.assinaturas, d.funcionarios)
    pdf.tabela(pd.DataFrame({
        "Assinatura": cpa["Assinatura"], "Status": cpa["Status"], "Qtd.": cpa["Qtd. usuários"],
        "Usuários": cpa["Usuários"].replace("", "-"),
        "Mensal": cpa["Gasto mensal"].map(calc.formatar_reais),
        "Por usuário": cpa["Custo por usuário"].map(calc.formatar_reais),
    }), larguras=(38, 20, 12, 140, 30, 30),
        alinhar=("LEFT", "LEFT", "CENTER", "LEFT", "RIGHT", "RIGHT"))
    return _bytes(pdf)


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
FORMATO_REAIS = '"R$" #,##0.00'
FORMATO_DATA = "DD/MM/YYYY"


def sanitizar_nome_aba(titulo: str, existentes: set[str] | list[str] = ()) -> str:
    """Sanitiza e deduplica o nome de uma aba para o limite de 31 caracteres do Excel.

    Remove caracteres proibidos ([]:*?/\\) e garante unicidade adicionando sufixo numérico.
    """
    base = "".join(c for c in str(titulo or "").strip() if c not in '[]:*?/\\')[:31].strip() or "Aba"
    nome = base
    contador = 1
    existentes_set = set(existentes)
    while nome in existentes_set:
        sufixo = f"_{contador}"
        nome = f"{base[:31 - len(sufixo)]}{sufixo}"
        contador += 1
    return nome


def _aba_excel(wb: Workbook, titulo: str, df: pd.DataFrame, moedas=(), datas=()):
    nome = sanitizar_nome_aba(titulo, wb.sheetnames)
    ws = wb.create_sheet(nome)
    cab_fill = PatternFill("solid", fgColor="2563EB")
    borda = Border(bottom=Side(style="thin", color="E2E8F0"))
    ws.append(list(df.columns))
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = cab_fill
        c.alignment = Alignment(vertical="center")
    for _, r in df.iterrows():
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in r.tolist()])
    for j, col in enumerate(df.columns, start=1):
        letra = get_column_letter(j)
        largura = max([len(str(col))] + [len(str(v)) for v in df[col].tolist()[:500]]) + 2
        ws.column_dimensions[letra].width = min(max(largura, 10), 60)
        for i in range(2, len(df) + 2):
            cel = ws.cell(row=i, column=j)
            cel.border = borda
            if col in moedas:
                cel.number_format = FORMATO_REAIS
            if col in datas:
                cel.number_format = FORMATO_DATA
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions
    return ws


def excel(abas: dict[str, pd.DataFrame], moedas=(), datas=()) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for titulo, df in abas.items():
        _aba_excel(wb, titulo, df, moedas, datas)
    if not wb.sheetnames:
        wb.create_sheet("Vazio")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _excel_assinaturas(df: pd.DataFrame, usuarios: dict) -> pd.DataFrame:
    return pd.DataFrame({
        "Assinatura": df["Nome"], "Categoria": df["Categoria"], "Plano": df["Plano"],
        "Valor (R$)": df["valor_num"], "Custo mensal (R$)": df["mensal"].round(2),
        "Vencimento": df["venc_data"], "Status": df["Status"], "Alerta": df["texto_alerta"],
        "Usuários": df["Nome"].map(lambda n: len(usuarios.get(str(n), []))),
    })


def por_categoria_xlsx(d: Dados) -> bytes:
    """RF-28: uma aba por categoria + uma aba de resumo."""
    df = d.assinaturas
    usuarios = calc.usuarios_por_assinatura(d.funcionarios)
    moedas = ("Valor (R$)", "Custo mensal (R$)", "Gasto mensal (R$)", "Gasto anual (R$)")
    abas = {}
    if not df.empty:
        cats = df.assign(Categoria=df["Categoria"].replace("", "Sem categoria"))
        resumo = cats.groupby("Categoria", as_index=False).agg(
            Assinaturas=("ID", "count"), **{"Gasto mensal (R$)": ("mensal", "sum")})
        resumo["Gasto anual (R$)"] = resumo["Gasto mensal (R$)"] * 12
        abas["Resumo"] = resumo.sort_values("Gasto mensal (R$)", ascending=False)
        for cat, grupo in cats.groupby("Categoria"):
            abas[str(cat)] = _excel_assinaturas(grupo.sort_values("Nome"), usuarios)
    else:
        abas["Resumo"] = pd.DataFrame(columns=["Categoria", "Assinaturas", "Gasto mensal (R$)"])
    return excel(abas, moedas=moedas, datas=("Vencimento",))


def colaboradores_xlsx(d: Dados) -> bytes:
    """RF-30 (Excel)."""
    cpa = calc.colaboradores_por_assinatura(d.assinaturas, d.funcionarios).rename(
        columns={"Gasto mensal": "Gasto mensal (R$)", "Custo por usuário": "Custo por usuário (R$)"})
    dep = calc.gasto_por_departamento(d.assinaturas, d.funcionarios).rename(
        columns={"Gasto mensal": "Gasto mensal (R$)", "Gasto anual": "Gasto anual (R$)"})
    pessoas = calc.custo_por_pessoa(d.assinaturas, d.funcionarios).drop(columns="ID").rename(
        columns={"Custo mensal": "Custo mensal (R$)"})
    return excel(
        {"Colaboradores por assinatura": cpa, "Gastos por departamento": dep,
         "Custo por pessoa": pessoas.sort_values("Custo mensal (R$)", ascending=False)},
        moedas=("Gasto mensal (R$)", "Custo por usuário (R$)", "Gasto anual (R$)", "Custo mensal (R$)"),
    )


def tabela_xlsx(df: pd.DataFrame, titulo: str, moedas=(), datas=()) -> bytes:
    """Botão 'Baixar Excel' embaixo de cada lista."""
    return excel({titulo: df}, moedas=moedas, datas=datas)
