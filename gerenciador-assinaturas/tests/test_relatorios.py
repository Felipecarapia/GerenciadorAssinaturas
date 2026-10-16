"""Testes dos relatórios em PDF e Excel: formatação, sanitização e integridade."""
from datetime import datetime
import pandas as pd
import pytest

from servicos import calculos as calc
from servicos import relatorios as rel


def _dados_teste():
    cols_ass = ["ID", "Nome", "Logo (URL)", "Categoria", "Plano", "Valor (R$)", "Vencimento", "Status", "Cadastrado em"]
    ass = pd.DataFrame([
        ("A1", "Figma", "", "Design", "Mensal", 150.0, "2026-11-01", "Ativa", "2025-01-01"),
        ("A2", "ChatGPT", "", "Produtividade", "Anual", 1200.0, "2026-10-20", "Ativa", "2025-01-01"),
        ("A3", "Claude", "", "Design/UI:Especial", "Mensal", 100.0, "2026-12-15", "Ativa", "2025-01-01"),
    ], columns=cols_ass)
    cols_func = ["ID", "Nome completo", "E-mail", "Departamento", "Cargo", "Assinaturas", "Cadastrado em"]
    func = pd.DataFrame([
        ("F1", "Ana Souza", "ana@empresa.com.br", "TI", "Engenheira", "Figma, ChatGPT", "2025-01-01"),
        ("F2", "Bia Lima", "bia@empresa.com.br", "Design", "Designer", "Figma, Claude", "2025-01-01"),
    ], columns=cols_func)
    hist = pd.DataFrame([
        {"Data e hora": "2026-10-15 10:00:00", "Origem": "app", "Quem": "admin", "Assinatura": "Figma",
         "Evento": "Valor alterado", "Campo": "Valor (R$)", "Valor anterior": "100.00", "Valor novo": "150.00"}
    ])
    momento = datetime(2026, 10, 16, 12, 0, 0)
    ass_prep = calc.preparar_assinaturas(ass, momento.date())
    return rel.Dados(ass_prep, func, hist, momento)


def test_sanitizar_nome_aba():
    # Remove caracteres inválidos no Excel: []:*?/\
    assert rel.sanitizar_nome_aba("Design/UI:Especial") == "DesignUIEspecial"
    # Trunca para 31 caracteres
    longo = "A" * 40
    assert len(rel.sanitizar_nome_aba(longo)) == 31
    # Deduplica em caso de colisão
    existentes = ["Resumo", "Design", "Design_1"]
    assert rel.sanitizar_nome_aba("Design", existentes) == "Design_2"
    # Deduplica com truncamento preservando o sufixo dentro de 31 caracteres
    assert len(rel.sanitizar_nome_aba("B" * 31, ["B" * 31])) == 31


def test_geracao_excel_por_categoria():
    dados = _dados_teste()
    conteudo = rel.por_categoria_xlsx(dados)
    assert isinstance(conteudo, bytes)
    assert len(conteudo) > 0
    # Verifica que o workbook pode ser lido pelo openpyxl
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(conteudo))
    assert "Resumo" in wb.sheetnames
    assert len(wb.sheetnames) >= 2


def test_geracao_excel_colaboradores():
    dados = _dados_teste()
    conteudo = rel.colaboradores_xlsx(dados)
    assert isinstance(conteudo, bytes)
    assert len(conteudo) > 0


def test_geracao_pdf_colaboradores_e_departamentos():
    dados = _dados_teste()
    pdf_bytes = rel.colaboradores_pdf(dados)
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF")


def test_geracao_pdf_a_vencer():
    dados = _dados_teste()
    pdf_bytes = rel.a_vencer_pdf(dados, 30)
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF")
